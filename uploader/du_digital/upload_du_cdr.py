#!/usr/bin/env python3
"""
Direct-to-database loader for DU Digital's CDR (Export Calls Report) into
du_cdr_daily_actual -- mirrors backend/src/modules/bulk-upload/
du-cdr-bulk.service.ts's importBatch() exactly. See upload_du_apr.py's own
header for the shared architecture (direct insert + upload_batch bookkeeping,
du_common.py for the ported parsing).

Unlike APR, no --force idempotency gate is needed: du_cdr_daily_actual's
UNIQUE KEY is (process_id, dashboard_label, uniqueid) -- uniqueid is
vicidial's own real per-call identifier, not a volatile per-batch value, so
re-running the same file (or an overlapping date range) always safely
upserts each real call to itself rather than duplicating it.

Reads the raw "CDR Raw" export's own column names directly (uniqueid,
call_date, status, user, full_name, campaign_id, user_group, status_name,
length_in_sec, queue_time, entry_date) -- no column renaming needed, since
the backend importer was deliberately built to accept this raw shape as-is.

Usage:
    py upload_du_cdr.py <file.xlsx|csv|txt> --country KOREA|THAILAND [--dry-run]
"""
from __future__ import annotations

import argparse
import sys
import uuid
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

sys.path.insert(0, str(Path(__file__).resolve().parent))
from du_common import (  # noqa: E402
    connect, get_process_id, load_db_config, parse_call_datetime, parse_count,
    write_staged_rows, write_upload_batch,
)

UPLOAD_TYPE_CODE = {"KOREA": "DU_CDR_KOREA", "THAILAND": "DU_CDR_THAILAND"}


def normalize_key(k: str) -> str:
    return "".join(ch for ch in k.lower() if ch.isalnum())


def read_rows(path: Path) -> list[dict[str, Any]]:
    lower = path.suffix.lower()
    if lower in (".csv", ".txt"):
        import csv
        with path.open(encoding="utf-8-sig", newline="") as f:
            sample = f.read(4096)
            f.seek(0)
            delim = "\t" if sample.count("\t") > sample.count(",") else ","
            return list(csv.DictReader(f, delimiter=delim))
    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]
    rows_iter = ws.iter_rows(values_only=True)
    header = [str(h) if h is not None else "" for h in next(rows_iter)]
    out = []
    for raw in rows_iter:
        if raw is None or all(v is None for v in raw):
            continue
        out.append({header[i]: raw[i] for i in range(len(header)) if i < len(raw)})
    wb.close()
    return out


def mapped_value(record: dict[str, Any], *candidates: str) -> Any:
    """Keeps the FIRST occurrence on a normalized-key collision -- see
    upload_du_apr.py's own mapped_value doc for the real WAIT/WAIT % bug this guards
    against; applied here too defensively even though CDR's own headers don't
    currently have an obvious colliding pair."""
    norm: dict[str, Any] = {}
    for k, v in record.items():
        nk = normalize_key(k)
        if nk not in norm:
            norm[nk] = v
    for c in candidates:
        v = norm.get(normalize_key(c))
        if v is not None and str(v).strip() != "":
            return v
    return None


def main() -> None:
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("file", type=Path)
    ap.add_argument("--country", required=True, choices=["KOREA", "THAILAND"])
    ap.add_argument("--dry-run", action="store_true")
    args = ap.parse_args()

    if not args.file.exists():
        sys.exit(f"File not found: {args.file}")

    print(f"Reading {args.file} ...")
    records = read_rows(args.file)
    print(f"  {len(records)} data row(s) read.")

    cfg = load_db_config()
    process_id = get_process_id(cfg)

    rows: list[dict[str, Any]] = []
    staged: list[tuple[int, dict[str, Any], str, str | None]] = []
    for i, rec in enumerate(records, start=2):
        uniqueid = str(mapped_value(rec, "uniqueid") or "").strip()
        if not uniqueid:
            staged.append((i, rec, "error", f'Row {i}: "uniqueid" is required'))
            continue
        call_date, call_datetime, hour = parse_call_datetime(mapped_value(rec, "call_date"))
        if not call_date:
            call_date, call_datetime, hour = parse_call_datetime(mapped_value(rec, "entry_date"))
        if not call_date:
            staged.append((i, rec, "error", f'Row {i}: "call_date" is required and could not be read'))
            continue
        status = str(mapped_value(rec, "status") or "").strip()
        if not status:
            staged.append((i, rec, "error", f'Row {i}: "status" is required'))
            continue
        rows.append({
            "uniqueid": uniqueid, "call_date": call_date.isoformat(), "call_datetime": call_datetime,
            "agent_user": (str(mapped_value(rec, "user") or "").strip() or None),
            "agent_name": (str(mapped_value(rec, "full_name", "Agent Name") or "").strip() or None),
            "phone_number": (str(mapped_value(rec, "phone_number_dialed", "phone_number") or "").strip() or None),
            "campaign_id": (str(mapped_value(rec, "campaign_id") or "").strip() or None),
            "user_group": (str(mapped_value(rec, "user_group") or "").strip() or None),
            "status": status,
            "status_name": (str(mapped_value(rec, "status_name") or "").strip() or None),
            "length_in_sec": parse_count(mapped_value(rec, "length_in_sec")),
            "queue_time": parse_count(mapped_value(rec, "queue_time")),
            "hour": hour,
        })
        staged.append((i, rec, "valid", None))

    print(f"  {len(rows)} row(s) valid, {len(staged) - len(rows)} skipped.")

    if args.dry_run:
        print("\n--dry-run: nothing written. First 3 mapped rows:")
        for r in rows[:3]:
            print(r)
        return

    imported = 0
    conn = connect(cfg)
    try:
        with conn.cursor() as cur:
            chunk_size = 1000
            for i in range(0, len(rows), chunk_size):
                chunk = rows[i:i + chunk_size]
                cur.executemany(
                    """INSERT INTO db_masmis.du_cdr_daily_actual
                         (id, process_id, dashboard_label, uniqueid, call_date, call_datetime, agent_user, agent_name,
                          phone_number, campaign_id, user_group, status, status_name, length_in_sec, queue_time_sec,
                          hour_of_day, data_source, source_reference, created_by)
                       VALUES (%s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, 'bulk_upload', %s, NULL)
                       ON DUPLICATE KEY UPDATE
                         call_date = VALUES(call_date), call_datetime = VALUES(call_datetime),
                         agent_user = VALUES(agent_user), agent_name = VALUES(agent_name),
                         phone_number = VALUES(phone_number), campaign_id = VALUES(campaign_id),
                         user_group = VALUES(user_group), status = VALUES(status), status_name = VALUES(status_name),
                         length_in_sec = VALUES(length_in_sec), queue_time_sec = VALUES(queue_time_sec),
                         hour_of_day = VALUES(hour_of_day)""",
                    [
                        (str(uuid.uuid4()), process_id, args.country, r["uniqueid"], r["call_date"], r["call_datetime"],
                         r["agent_user"], r["agent_name"], r["phone_number"], r["campaign_id"], r["user_group"],
                         r["status"], r["status_name"], r["length_in_sec"], r["queue_time"], r["hour"],
                         "du_digital_auto")
                        for r in chunk
                    ],
                )
                imported += len(chunk)
                print(f"  {imported}/{len(rows)} rows upserted", end="\r")
        conn.commit()
        print()
    finally:
        conn.close()

    print(f"Imported {imported}/{len(rows)} row(s) into du_cdr_daily_actual ({args.country}).")

    # The rows above are already committed. Recent Uploads bookkeeping is best-effort: if db_masmis
    # lacks upload_batch the import still succeeded, so do not fail the run (the auto-import job
    # treats a non-zero exit as a failed import).
    try:
        batch_id = write_upload_batch(
            cfg, UPLOAD_TYPE_CODE[args.country], "du_cdr_daily_actual", args.file.name, args.file.stat().st_size,
            len(records), len(rows), imported, len(staged) - len(rows),
        )
        write_staged_rows(cfg, batch_id, staged)
        print(f"Recorded upload_batch {batch_id} -- will show in the Uploader page's Recent Uploads.")
    except Exception as exc:  # noqa: BLE001 -- bookkeeping only; the data import already succeeded
        print(f"Warning: rows imported but upload_batch was not recorded ({exc}).")


if __name__ == "__main__":
    main()
